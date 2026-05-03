import openpyxl

wb = openpyxl.load_workbook("../FS.xlsx", data_only=False)
for sheet in wb.sheetnames:
    ws = wb[sheet]
    print(f"\n--- Sheet: {sheet} ---")
    count = 0
    for row in ws.iter_rows(min_row=1, max_row=40):
        for cell in row:
            if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                print(f"Cell {cell.coordinate}: {cell.value}")
                count += 1
                if count > 15:
                    break
        if count > 15:
            break
