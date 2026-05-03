import openpyxl

wb = openpyxl.load_workbook("../FS.xlsx", data_only=False)
for sheet in wb.sheetnames:
    ws = wb[sheet]
    print(f"\n--- Sheet: {sheet} ---")
    for row in ws.iter_rows(min_row=1, max_row=25):
        row_data = []
        for cell in row[:6]: # A to F
            val = cell.value
            if val is not None:
                if isinstance(val, str) and val.strip() == "":
                    row_data.append("EMPTY")
                else:
                    row_data.append(str(val))
            else:
                row_data.append("None")
        if any(v != "None" and v != "EMPTY" for v in row_data):
            print(f"Row {row[0].row}: {row_data}")
